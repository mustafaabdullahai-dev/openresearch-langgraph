"""Tests for document extraction, upload security, and vector-store RAG."""

from __future__ import annotations

import asyncio

import pytest
from app.config.settings import settings
from app.services.documents.extractor import extract_text
from app.services.security import (
    MAX_UPLOAD_SIZE_BYTES,
    is_allowed_upload,
)
from app.services.vectorstore.chroma_store import (
    ChromaVectorStore,
    _LocalEmbeddingFunction,
)

_TXT = "Solar panels degrade by roughly 0.5% per year on average.\n" * 40


def test_extract_txt(tmp_path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text(_TXT)
    assert extract_text(str(path)) == _TXT


def test_extract_markdown(tmp_path) -> None:
    path = tmp_path / "notes.md"
    body = "# Title\n\nSome **bold** content."
    path.write_text(body)
    assert extract_text(str(path)) == body


def test_extract_docx(tmp_path) -> None:
    from docx import Document

    path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("First paragraph of the report.")
    doc.add_paragraph("Second paragraph.")
    doc.save(path)
    text = extract_text(str(path))
    assert "First paragraph of the report." in text
    assert "Second paragraph." in text


def test_extract_unsupported_extension_raises(tmp_path) -> None:
    path = tmp_path / "note.exe"
    path.write_bytes(b"\x00\x01")
    with pytest.raises(ValueError):
        extract_text(str(path))


def test_is_allowed_upload_valid_txt() -> None:
    allowed, reason = is_allowed_upload("notes.txt", 1024)
    assert allowed is True
    assert reason == ""


def test_is_allowed_upload_rejects_unknown_extension() -> None:
    allowed, reason = is_allowed_upload("notes.exe", 1024)
    assert allowed is False
    assert "not supported" in reason


def test_is_allowed_upload_rejects_oversize() -> None:
    allowed, reason = is_allowed_upload("notes.txt", MAX_UPLOAD_SIZE_BYTES + 1)
    assert allowed is False
    assert "exceeds" in reason


def test_is_allowed_upload_case_insensitive() -> None:
    allowed, _reason = is_allowed_upload("notes.PDF", 10)
    assert allowed is True


def test_hash_embedding_is_deterministic() -> None:
    fn = _LocalEmbeddingFunction()
    first = [list(v) for v in fn(["hello world", "solar panels"])]
    second = [list(v) for v in fn(["hello world", "solar panels"])]
    assert first == second
    assert len(first[0]) == 128


def test_hash_embedding_distinguishes_text() -> None:
    fn = _LocalEmbeddingFunction()
    a = [list(v) for v in fn(["solar panels"])]
    b = [list(v) for v in fn(["tomato soup"])]
    assert a[0] != b[0]


def test_hash_embedding_vectors_are_unit_norm() -> None:
    fn = _LocalEmbeddingFunction()
    vectors = [list(v) for v in fn(["solar panels", "quantum computing"])]
    for vec in vectors:
        norm = float(sum(v * v for v in vec) ** 0.5)
        assert abs(norm - 1.0) < 1e-6


def test_vector_store_add_and_similarity_roundtrip(
    monkeypatch: pytest.MonkeyPatch, tmp_path
) -> None:
    monkeypatch.setattr(settings, "VECTOR_DB_PATH", str(tmp_path / "vec_roundtrip"))
    store = ChromaVectorStore()

    async def _work() -> None:
        await store.add_documents(
            ids=["doc-1", "doc-2"],
            documents=[
                "Solar panels degrade slowly with time and temperature.",
                "Tomatoes grow best in warm climates with rich soil.",
            ],
            metadatas=[
                {"filename": "solar.txt", "chunk_index": 0},
                {"filename": "garden.txt", "chunk_index": 0},
            ],
        )
        chunks = await store.similarity_search("solar degradation", k=2)
        assert chunks, "expected at least one retrieved chunk"
        top = chunks[0]
        assert "solar" in top.content.lower()
        assert top.source_filename == "solar.txt"
        assert 0.0 <= top.score <= 1.0

    asyncio.run(_work())


def test_vector_store_persists_across_instances(
    monkeypatch: pytest.MonkeyPatch, tmp_path
) -> None:
    monkeypatch.setattr(settings, "VECTOR_DB_PATH", str(tmp_path / "vec_persist"))
    store = ChromaVectorStore()

    async def _seed_and_read() -> None:
        await store.add_documents(
            ids=["doc-1"],
            documents=["Quantum computing uses qubits for computation."],
            metadatas=[{"filename": "quantum.txt", "chunk_index": 0}],
        )

    asyncio.run(_seed_and_read())

    second = ChromaVectorStore()

    async def _read() -> None:
        chunks = await second.similarity_search("qubits", k=1)
        assert chunks
        assert "quantum" in chunks[0].content.lower()

    asyncio.run(_read())
